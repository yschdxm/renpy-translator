"""翻译文本导出：对话 / UI 字符串 / 人名 / 术语 -> txt / json / xlsx / docx / 翻译Word

纯逻辑无 UI：server 层在 db_call 里调 collect_data 取数，
在 run_sync 里调 write_export_file 写盘（openpyxl / python-docx
惰性导入，首次导出对应格式才加载——打包 spec 惰性清单已登记）。

docx / xlsx 字体约定：中文宋体、西文 Times New Roman（见 _set_docx_fonts
与 _font_pair）。翻译Word（tword）：只写原文、每条一段，供译者翻译后导回。
"""

import json
import re
from datetime import datetime
from pathlib import Path

EXPORT_TYPES = ('dialogue', 'ui', 'names', 'glossary')
EXPORT_FORMATS = ('txt', 'json', 'xlsx', 'docx', 'tword')

# 中英文字体约定（docx / xlsx 共用）
_EN_FONT = 'Times New Roman'
_CN_FONT = '宋体'
# CJK 统一表意文字 + 扩展A + 兼容表意 + 中文标点 + 全角字符
_CJK_RE = re.compile(
    r'[㐀-䶿一-鿿豈-﫿　-〿＀-￯]')

# 分节标题（txt 节名 / xlsx sheet 名 / docx 标题）
TYPE_TITLES = {
    'dialogue': '对话翻译',
    'ui': 'UI 字符串',
    'names': '人名表',
    'glossary': '术语表',
}

# json 输出的顶层键（英文、稳定，程序友好）
SECTION_KEYS = {
    'dialogue': 'dialogues',
    'ui': 'ui_texts',
    'names': 'names',
    'glossary': 'glossary',
}

# 列定义单一事实源：type -> [(列key, 数据字段, 中文表头)]，顺序即默认列顺序
COLUMN_DEFS = {
    'dialogue': [
        ('original', 'original_text', '原文'),
        ('translated', 'translated_text', '译文'),
        ('character', 'character', '说话人'),
        ('file', 'file_path', '来源文件'),
        ('line', 'line_number', '行号'),
        ('label', 'label', '标签'),
        ('status', 'is_translated', '已翻译'),
    ],
    'ui': [
        ('original', 'original_text', '原文'),
        ('translated', 'translated_text', '译文'),
        ('file', 'file_path', '来源文件'),
        ('line', 'line_number', '行号'),
        ('label', 'label', '标签'),
        ('context', 'context_hint', '上下文提示'),
        ('status', 'is_translated', '已翻译'),
    ],
    'names': [
        ('original', 'display_name', '原名'),
        ('translated', 'cn_name', '译名'),
        ('variable', 'variable', '变量名'),
        ('lines', 'lines_count', '台词数'),
    ],
    'glossary': [
        ('en', 'en_term', '英文术语'),
        ('cn', 'cn_term', '中文术语'),
        ('type', 'term_type', '类型'),
        ('source', 'source', '来源'),
    ],
}

# txt 输出时并入「[...]」头部行的列（其余列逐行「表头：值」）
_TXT_HEADER_KEYS = {
    'dialogue': ('file', 'line', 'character', 'label'),
    'ui': ('file', 'line', 'label'),
}


def resolve_columns(types: list[str], columns: dict | None) -> dict[str, list[str]]:
    """校验并解析列选择：columns 的 key 必须是 types 子集，列名走白名单，
    未指定的类型取全部列；结果顺序对齐 COLUMN_DEFS。"""
    columns = columns or {}
    extra = set(columns) - set(types)
    if extra:
        raise ValueError(f'导出列包含未选中的类型: {sorted(extra)}')
    result = {}
    for t in types:
        defs = COLUMN_DEFS[t]
        sel = columns.get(t)
        if sel is None:
            result[t] = [k for k, _, _ in defs]
            continue
        valid = {k for k, _, _ in defs}
        bad = [c for c in sel if c not in valid]
        if bad:
            raise ValueError(f'{TYPE_TITLES[t]} 未知导出列: {bad}')
        if not sel:
            raise ValueError(f'{TYPE_TITLES[t]} 至少选择一列')
        chosen = set(sel)
        result[t] = [k for k, _, _ in defs if k in chosen]
    return result


def collect_data(db, types: list[str], columns: dict | None,
                 flt: dict | None = None) -> dict[str, list[dict]]:
    """按类型取行并按 columns 裁剪（行键为英文列 key）。

    flt（仅作用于其 content_type 指定的 dialogue/ui）：
    {content_type, filter_mode, search, character}
    """
    cols = resolve_columns(types, columns)
    data = {}
    # 规范类型序（与前端勾选顺序无关）：tword 导出/导入按序对齐依赖此顺序
    for t in EXPORT_TYPES:
        if t not in types:
            continue
        if t == 'dialogue':
            if flt and flt.get('content_type') == 'dialogue':
                rows = db.get_dialogues_filtered(
                    flt.get('filter_mode', 'all'), flt.get('character', ''),
                    flt.get('search', ''))
            else:
                rows = db.get_all_dialogues()
        elif t == 'ui':
            if flt and flt.get('content_type') == 'ui':
                rows = db.get_ui_texts_filtered(
                    flt.get('filter_mode', 'all'), flt.get('search', ''))
            else:
                rows = db.get_all_ui_texts()
        elif t == 'names':
            rows = db.get_characters()
        else:
            rows = db.get_glossary_rows()
        data[t] = [_project_row(r, t, cols[t]) for r in rows]
    return data


def _project_row(row: dict, t: str, sel: list[str]) -> dict:
    defs = {k: f for k, f, _ in COLUMN_DEFS[t]}
    return {k: row.get(defs[k]) for k in sel}


def _headers(t: str, sel: list[str]) -> list[str]:
    header_of = {k: h for k, _, h in COLUMN_DEFS[t]}
    return [header_of[k] for k in sel]


def _fmt(value) -> str:
    """单元格文本化：布尔转 是/否，None 转空串"""
    if value is None:
        return ''
    if isinstance(value, bool):
        return '是' if value else '否'
    return str(value)


# ---- txt ----

def render_txt(data: dict[str, list[dict]], meta: dict) -> str:
    """分节纯文本（写盘用 utf-8-sig，Windows 记事本直接可读）"""
    out = [f"{meta['project']} 翻译文本导出",
           f"导出时间：{meta['exported_at']}", '']
    for t, rows in data.items():
        out.append(f"===== {TYPE_TITLES[t]}（共 {len(rows)} 条） =====")
        out.append('')
        header_keys = _TXT_HEADER_KEYS.get(t, ())
        for row in rows:
            head = _txt_head(row, t, header_keys)
            if head:
                out.append(head)
            for k, v in row.items():
                if k in header_keys:
                    continue
                label = _headers(t, [k])[0]
                out.append(f'{label}：{_fmt(v)}')
            out.append('')
    return '\n'.join(out)


def _txt_head(row: dict, t: str, header_keys: tuple) -> str:
    if not any(k in row for k in header_keys):
        return ''
    loc = ''
    if 'file' in row or 'line' in row:
        f, ln = _fmt(row.get('file')), _fmt(row.get('line'))
        loc = f'{f}:{ln}' if f and ln else (f or ln)
    name = _fmt(row.get('character')) if 'character' in row else ''
    label = _fmt(row.get('label')) if 'label' in row else ''
    parts = loc
    if label:
        parts += f' ({label})' if parts else f'({label})'
    if name:
        parts += f' {name}' if parts else name
    return f'[{parts}]' if parts else ''


# ---- json ----

def render_json(data: dict[str, list[dict]], meta: dict) -> str:
    doc = {'project': meta['project'], 'exported_at': meta['exported_at'],
           'format_version': 1}
    for t, rows in data.items():
        doc[SECTION_KEYS[t]] = rows
    return json.dumps(doc, ensure_ascii=False, indent=2)


# ---- xlsx ----

def _font_pair(text: str, bold: bool = False):
    """中文宋体 / 西文 Times New Roman：返回 (单元格值, 基础字体)。

    纯中文或纯西文用普通单元格 + 整格字体（大多数单元格，零额外开销）；
    中英混排用富文本按字符集分段（基础字体返回 None，由 run 自带字体）。
    """
    from openpyxl.styles import Font

    if not text or not _CJK_RE.search(text):
        return text, Font(name=_EN_FONT, bold=bold)
    if all(_CJK_RE.match(c) or c.isspace() for c in text):
        return text, Font(name=_CN_FONT, bold=bold)
    return _rich_text(text, bold), None


def _rich_text(text: str, bold: bool):
    """中英混排富文本：CJK 段宋体、其余段 Times New Roman"""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    cn = InlineFont(rFont=_CN_FONT, b=bold)
    en = InlineFont(rFont=_EN_FONT, b=bold)
    blocks, buf, cur = [], '', None
    for ch in text:
        cjk = bool(_CJK_RE.match(ch))
        if cur is not None and cjk != cur:
            blocks.append(TextBlock(cn if cur else en, buf))
            buf = ''
        cur = cjk
        buf += ch
    if buf:
        blocks.append(TextBlock(cn if cur else en, buf))
    return CellRichText(*blocks)


def write_xlsx(path: Path, data: dict[str, list[dict]], meta: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    wrap = Alignment(wrap_text=True, vertical='top')
    for t, rows in data.items():
        ws = wb.create_sheet(TYPE_TITLES[t])
        sel = list(rows[0].keys()) if rows else [k for k, _, _ in COLUMN_DEFS[t]]
        for j, h in enumerate(_headers(t, sel), start=1):
            v, f = _font_pair(h, bold=True)
            c = ws.cell(row=1, column=j, value=v)
            if f:
                c.font = f
        for i, row in enumerate(rows, start=2):
            for j, k in enumerate(sel, start=1):
                v, f = _font_pair(_fmt(row.get(k)))
                c = ws.cell(row=i, column=j, value=v)
                if f:
                    c.font = f
                c.alignment = wrap
        for j, k in enumerate(sel, start=1):
            ws.column_dimensions[get_column_letter(j)].width = (
                60 if k in ('original', 'translated', 'context') else 18)
        ws.freeze_panes = 'A2'
    wb.save(path)


# ---- docx ----

def write_docx(path: Path, data: dict[str, list[dict]], meta: dict) -> None:
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn

    doc = Document()
    _set_docx_fonts(doc)
    doc.add_heading(f"{meta['project']} 翻译文本导出", level=0)
    doc.add_paragraph(f"导出时间：{meta['exported_at']}")
    body = doc._body._element
    sect_pr = body.find(qn('w:sectPr'))
    for t, rows in data.items():
        doc.add_heading(f"{TYPE_TITLES[t]}（共 {len(rows)} 条）", level=1)
        sel = list(rows[0].keys()) if rows else [k for k, _, _ in COLUMN_DEFS[t]]
        headers = _headers(t, sel)
        # lxml 把解析好的大子树插进文档是超线性慢（22k 行整表插入 13s），
        # 分块成小表插入可降到 ~1.4s；同列宽相邻表格在 Word 中视觉连续
        for i in range(0, max(len(rows), 1), _DOCX_CHUNK_ROWS):
            chunk = rows[i:i + _DOCX_CHUNK_ROWS]
            tbl = parse_xml(_table_xml(
                headers if i == 0 else [],
                [[_fmt(r.get(k)) for k in sel] for r in chunk]))
            # sectPr 必须是 body 最后一个子元素，表插到它前面
            if sect_pr is not None:
                sect_pr.addprevious(tbl)
            else:
                body.append(tbl)
    doc.save(path)


def _set_docx_fonts(doc) -> None:
    """中文宋体 / 西文 Times New Roman：统一设置正文与标题样式

    eastAsia 控制中文字形，ascii/hAnsi 控制西文；主题字体属性存在时
    优先于显式字体，必须清除。
    """
    from docx.oxml.ns import qn

    for name in ('Normal', 'Title', 'Heading 1'):
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        style.font.name = _EN_FONT
        r_fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            r_fonts.attrib.pop(qn(f'w:{attr}'), None)
        r_fonts.set(qn('w:eastAsia'), _CN_FONT)


_DOCX_CHUNK_ROWS = 2000


def _table_xml(headers: list[str], rows: list[list[str]]) -> str:
    """整表 XML 一次性解析：python-docx 逐行 add_row 在大表下 O(n²) 实测
    要 95s+，拼字符串 + parse_xml 约 1s。headers 为空时不输出表头行。"""
    from xml.sax.saxutils import escape

    from docx.oxml.ns import nsdecls

    def cell(text: str) -> str:
        t = escape(text).replace(
            '\n', '</w:t><w:br/><w:t xml:space="preserve">')
        return ('<w:tc><w:tcPr><w:tcW w:w="0" w:type="auto"/></w:tcPr>'
                f'<w:p><w:r><w:t xml:space="preserve">{t}</w:t></w:r></w:p>'
                '</w:tc>')

    parts = [f'<w:tbl {nsdecls("w")}><w:tblPr><w:tblStyle w:val="TableGrid"/>'
             '<w:tblW w:w="0" w:type="auto"/></w:tblPr>']
    # 列数：无表头的后续分块从数据行推断，保证 tblGrid 与单元格数一致
    n_cols = len(headers) or (len(rows[0]) if rows else 0)
    parts.append('<w:tblGrid>' + '<w:gridCol/>' * n_cols + '</w:tblGrid>')
    if headers:
        parts.append('<w:tr>' + ''.join(cell(h) for h in headers) + '</w:tr>')
    for row in rows:
        parts.append('<w:tr>' + ''.join(cell(v) for v in row) + '</w:tr>')
    parts.append('</w:tbl>')
    return ''.join(parts)


# ---- 翻译 Word（tword：仅原文，每条一段） ----

def write_tword(path: Path, originals: list[str]) -> None:
    """翻译 Word：只写原文、每条一段，除此之外什么都没有（供译者翻译后导回，
    导入按段落顺序与条目一一对应）。段内换行用 <w:br/>，保证 1 条 = 1 段。"""
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn

    doc = Document()
    _set_docx_fonts(doc)
    body = doc._body._element
    sect_pr = body.find(qn('w:sectPr'))
    for i in range(0, len(originals), _DOCX_CHUNK_ROWS):
        root = parse_xml(_paras_xml(originals[i:i + _DOCX_CHUNK_ROWS]))
        for p in list(root):
            if sect_pr is not None:
                sect_pr.addprevious(p)
            else:
                body.append(p)
    doc.save(path)


def _paras_xml(lines: list[str]) -> str:
    """段落序列包一层根元素便于一次 parse_xml（避免逐段 add_paragraph）"""
    from xml.sax.saxutils import escape

    from docx.oxml.ns import nsdecls

    parts = [f'<root {nsdecls("w")}>']
    for ln in lines:
        t = escape(ln).replace(
            '\n', '</w:t><w:br/><w:t xml:space="preserve">')
        parts.append(f'<w:p><w:r><w:t xml:space="preserve">{t}</w:t></w:r></w:p>')
    parts.append('</root>')
    return ''.join(parts)


# ---- 顶层编排 ----

def write_export_file(data: dict[str, list[dict]], exports_dir: Path,
                      project_name: str, fmt: str,
                      filename: str = '') -> tuple[Path, dict[str, int]]:
    """按格式写导出文件，返回 (文件路径, 各类型条数)。"""
    if fmt not in EXPORT_FORMATS:
        raise ValueError(f'未知导出格式: {fmt}')
    meta = {'project': project_name,
            'exported_at': datetime.now().isoformat(timespec='seconds')}
    path = _unique_path(exports_dir, _sanitize(filename), project_name, fmt)
    if fmt == 'txt':
        path.write_text(render_txt(data, meta), encoding='utf-8-sig')
    elif fmt == 'json':
        path.write_text(render_json(data, meta), encoding='utf-8')
    elif fmt == 'xlsx':
        write_xlsx(path, data, meta)
    elif fmt == 'tword':
        # 规范类型序 + 只写原文（server 层已把 columns 裁为 original）
        originals = [str(row.get('original') or '')
                     for t in EXPORT_TYPES if t in data
                     for row in data[t]]
        write_tword(path, originals)
    else:
        write_docx(path, data, meta)
    return path, {t: len(rows) for t, rows in data.items()}


def _sanitize(filename: str) -> str:
    """用户自定义主名：与 _exports_dir 同款字符白名单，去扩展名"""
    name = Path(filename).stem if filename else ''
    return ''.join(c for c in name if c.isalnum() or c in '._- ').strip()


def _unique_path(exports_dir: Path, base: str, project_name: str,
                 fmt: str) -> Path:
    """空名自动生成带时间戳文件名；指定名冲突时追加 -2/-3，不静默覆盖"""
    ext = 'docx' if fmt == 'tword' else fmt
    if not base:
        base = f"{project_name}-texts-{datetime.now():%Y%m%d-%H%M%S}"
    path = exports_dir / f'{base}.{ext}'
    n = 2
    while path.exists():
        path = exports_dir / f'{base}-{n}.{ext}'
        n += 1
    return path
